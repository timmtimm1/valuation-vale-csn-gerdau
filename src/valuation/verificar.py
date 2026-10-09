"""Confere a planilha contra o modelo em Python.

A planilha sai do openpyxl só com fórmulas, sem valores. Aqui o LibreOffice abre
o arquivo, calcula tudo e grava uma cópia; lemos os valores dessa cópia e
comparamos com o que `modelo.py` calculou por conta própria. Se uma fórmula do
Excel divergir da conta em Python, aparece aqui.
"""

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from valuation import historico, modelo, premissas
from valuation.empresas import SAIDA

TOLERANCIA = 1e-6  # relativa
ERROS_EXCEL = ("#DIV/0!", "#REF!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "Err:")


def recalcular(planilha: Path, destino: Path) -> Path:
    """Abre no LibreOffice sem tela e salva com os valores calculados."""
    soffice = shutil.which("soffice")
    if soffice is None:
        raise RuntimeError("LibreOffice (soffice) não encontrado")
    subprocess.run(
        [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", destino, planilha],
        check=True,
        capture_output=True,
        timeout=180,
    )
    return destino / planilha.name


def _linha(ws: Worksheet, rotulo: str) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value == rotulo:
            return r
    raise KeyError(f"{ws.title}: linha '{rotulo}' não encontrada")


def _perto(a: float, b: float, tolerancia: float = TOLERANCIA) -> bool:
    return abs(a - b) <= tolerancia * max(1.0, abs(a), abs(b))


def conferir(ticker: str) -> list[str]:
    """Lista de divergências (vazia quando a planilha e o modelo concordam)."""
    dados = premissas.carregar(ticker)
    r = modelo.rodar(ticker, dados)
    proj, v = r["projecao"], r["valuation"]
    problemas: list[str] = []

    with tempfile.TemporaryDirectory() as pasta:
        calculada = recalcular(SAIDA / f"valuation_{ticker}.xlsx", Path(pasta))
        wb = load_workbook(calculada, data_only=True)

    def comparar(onde: str, obtido: Any, esperado: float, tol: float = TOLERANCIA) -> None:
        numero = isinstance(obtido, int | float)
        if not numero or not _perto(float(obtido), float(esperado), tol):
            problemas.append(f"{onde}: planilha {obtido!r}, modelo {esperado:.6f}")

    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str) and celula.value.startswith(ERROS_EXCEL):
                    problemas.append(f"{ws.title}!{celula.coordinate}: {celula.value}")

    # Projeção: todas as linhas do modelo, todos os anos
    ws = wb["Projeção"]
    rotulos = {
        "receita": "Receita líquida",
        "ebitda": "EBITDA",
        "custo_caixa": "(-) Custo dos produtos (sem depreciação)",
        "despesas": "(-) Despesas operacionais (vendas, administrativas e outras)",
        "ebit": "EBIT",
        "lucro_liquido": "Lucro líquido",
        "nopat": "NOPAT (EBIT sem equivalência, após imposto)",
        "capex": "Capex",
        "da": "(-) Depreciação e amortização",
        "variacao_giro": "Variação do capital de giro (investimento em giro)",
        "fcff": "FCFF - fluxo de caixa livre da empresa",
        "fcfe": "FCFE - fluxo de caixa livre do acionista",
        "caixa_total": "Caixa e aplicações",
        "patrimonio_liquido": "Patrimônio líquido",
        "checagem_balanco": "Checagem: ativo - passivo (tem de ser zero)",
        "roic": "ROIC (NOPAT / capital investido inicial)",
    }
    for nome, rotulo in rotulos.items():
        linha_xl = _linha(ws, rotulo)
        for i, ano in enumerate(proj.columns):
            comparar(f"Projeção {nome} {ano}", ws.cell(linha_xl, 4 + i).value, proj.loc[nome, ano])
    base_checagem = ws.cell(_linha(ws, rotulos["checagem_balanco"]), 3).value
    comparar("Projeção checagem do ano-base", base_checagem, 0.0)

    # DCF
    ws = wb["DCF"]
    for nome, rotulo in {
        "wacc": "WACC",
        "ev": "Valor da empresa (Enterprise Value)",
        "equity": "Valor do acionista (Equity Value)",
        "preco_justo": "Preço justo por ação",
        "preco_fcfe": "Preço por ação pelo FCFE",
        "peso_terminal": "Peso do valor terminal no valor da empresa",
    }.items():
        comparar(f"DCF {nome}", ws.cell(_linha(ws, rotulo), 4).value, v[nome])
    grade = modelo.sensibilidade(r["base"], proj, r["premissas"])
    topo = next(i for i in range(1, ws.max_row + 1) if ws.cell(i, 3).value == "WACC \\ g")
    for i in range(5):
        for j in range(5):
            comparar(
                f"Sensibilidade [{i},{j}]", ws.cell(topo + 1 + i, 4 + j).value, grade.iat[i, j]
            )

    # Histórico: indicadores em fórmula contra os calculados em pandas
    ws = wb["Histórico"]
    h = historico.carregar()
    h = h[h["ticker"] == ticker].reset_index(drop=True)
    for nome, rotulo in {
        "ebitda": "EBITDA (EBIT + depreciação)",
        "nopat": "NOPAT (lucro operacional após imposto de 34%)",
        "margem_ebitda": "Margem EBITDA",
        "margem_ebitda_recorrente": "Margem EBITDA sem perdas por recuperabilidade",
        "prazo_estoque": "Prazo médio de estoque (dias)",
        "divida_liquida_ebitda": "Dívida líquida / EBITDA",
        "roic": "ROIC (NOPAT / capital de giro + ativo fixo)",
    }.items():
        linha_xl = _linha(ws, rotulo)
        for i, periodo in enumerate(h["periodo"]):
            # o CSV guarda seis casas; a planilha calcula com os valores cheios
            obtido = ws.cell(linha_xl, 3 + i).value
            comparar(f"Histórico {nome} {periodo}", obtido, h.at[i, nome], 1e-5)

    # Cenários: a planilha viva bate com o comparativo e o caixa fecha
    ws = wb["Cenários"]
    linha_xl = _linha(ws, "Diferença para o comparativo acima (zero sem ajuste manual)")
    comparar("Cenários: vivo x comparativo", ws.cell(linha_xl, 4).value, 0.0)
    linha_xl = _linha(ws, "Checagem: bate com a variação de caixa da projeção (zero)")
    for i in range(len(proj.columns)):
        comparar(f"Cenários: caixa {proj.columns[i]}", ws.cell(linha_xl, 4 + i).value, 0.0)

    # Seletor e ajuste manual: mexer na planilha tem de dar o mesmo preço que o Python.
    for cenario, manual in SIMULACOES:
        esperado = modelo.rodar(ticker, dados, cenario, manual)["valuation"]["preco_justo"]
        obtido = _preco_simulado(ticker, cenario, manual)
        comparar(f"Simulação {cenario} {manual or ''}", obtido, esperado)
    return problemas


# Os dois cenários extremos, um ajuste manual dentro da faixa e um fora dela, que a
# planilha tem de prender no limite como o Python faz.
SIMULACOES: tuple[tuple[str, dict[str, float]], ...] = (
    ("pessimista", {}),
    ("otimista", {}),
    ("moderado", {"crescimento_receita": 0.02, "ajuste_wacc": 0.005}),
    ("pessimista", {"margem_ebitda": 0.99}),
)
NOMES = {"pessimista": "Pessimista", "moderado": "Moderado", "otimista": "Otimista"}
COLUNA_MANUAL = 8


def _preco_simulado(ticker: str, cenario: str, manual: dict[str, float]) -> Any:
    """Troca o cenário e os ajustes manuais numa cópia da planilha, recalcula e lê o preço."""
    wb = load_workbook(SAIDA / f"valuation_{ticker}.xlsx")
    ws = wb["Premissas"]
    ws["D4"].value = NOMES[cenario]
    for nome, valor in manual.items():
        ws.cell(_linha(ws, premissas.VARIAVEIS[nome]), COLUNA_MANUAL).value = valor
    with tempfile.TemporaryDirectory() as pasta:
        entrada = Path(pasta) / "entrada"
        entrada.mkdir()
        copia = entrada / f"valuation_{ticker}.xlsx"
        wb.save(copia)
        calculada = load_workbook(recalcular(copia, Path(pasta)), data_only=True)
    dcf = calculada["DCF"]
    return dcf.cell(_linha(dcf, "Preço justo por ação"), 4).value
