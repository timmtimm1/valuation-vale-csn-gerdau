"""A planilha em fórmulas tem de dar o mesmo resultado que o modelo em Python."""

import shutil

import pytest
from openpyxl import Workbook

from valuation import planilha, verificar
from valuation.empresas import FOCO

precisa_do_libreoffice = pytest.mark.skipif(
    shutil.which("soffice") is None, reason="precisa do LibreOffice"
)


@pytest.fixture(scope="module")
def molde() -> Workbook:
    return planilha.montar()


def test_abas_na_ordem_do_estudo(molde: Workbook) -> None:
    assert molde.sheetnames == planilha.ORDEM
    assert molde.sheetnames[0] == "Painel"


def test_toda_linha_do_modelo_tem_explicacao(molde: Workbook) -> None:
    # nas abas de conta, toda linha com número tem uma frase ao lado dizendo o que é
    for nome, primeira_coluna in (("Projeção", 3), ("FCFF", 4), ("WACC", 4), ("Valor justo", 4)):
        ws = molde[nome]
        coluna_nota = max(c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value)
        for r in range(5, ws.max_row + 1):
            tem_conta = any(
                isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.startswith("=")
                for c in range(primeira_coluna, coluna_nota)
            )
            rotulo = ws.cell(r, 2).value
            if tem_conta and rotulo and not str(rotulo).startswith("="):
                assert ws.cell(r, coluna_nota).value, f"{nome}: sem explicação em {rotulo!r}"


@precisa_do_libreoffice
@pytest.mark.parametrize("ticker", [e.ticker for e in FOCO])
def test_planilha_recalculada_bate_com_o_modelo(molde: Workbook, ticker: str) -> None:
    assert verificar.conferir(ticker, molde) == []
